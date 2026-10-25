import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.core.config import get_settings
from app.models.source import Source, SourceKind
from app.processors.base import ProcessingError, ProcessorContext
from app.processors.registry import process_source
from app.processors.documents import process_docx, process_pdf
from app.processors.media import process_audio, process_image, process_youtube_transcript, window_transcript
from app.processors.text import pack_paragraphs, split_markdown
from app.processors.web import html_to_markdown, _MarkdownExtractor
from app.providers.base import ImageDescription, OCRObservation, TranscriptionResult


def make_pdf(pages: list[str]) -> bytes:
    """A valid PDF with one Helvetica text block per page. An empty string makes a page with no text."""

    def escape(line: str) -> str:
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    page_ids = [4 + 2 * index for index in range(len(pages))]
    objects = {
        1: "<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{' '.join(f'{pid} 0 R' for pid in page_ids)}] /Count {len(pages)} >>",
        3: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for pid, text in zip(page_ids, pages):
        operations = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
        operations += [f"({escape(line)}) Tj T*" for line in text.split("\n")] if text else []
        stream = "\n".join(operations + ["ET"]) if text else ""
        objects[pid] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {pid + 1} 0 R >>"
        )
        objects[pid + 1] = f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream"

    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(output)
        output += f"{number} 0 obj\n{objects[number]}\nendobj\n".encode("latin-1")
    xref = len(output)
    size = max(objects) + 1
    output += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for number in range(1, size):
        output += f"{offsets[number]:010d} 00000 n \n".encode()
    output += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(output)


class FakeOCR:
    def __init__(self, text: str = "Scanned text from OCR", error: Exception | None = None) -> None:
        self.text, self.error, self.calls = text, error, 0

    def extract(self, frame_paths):
        self.calls += 1
        if self.error:
            raise self.error
        return [OCRObservation(frame_path=str(path), text=self.text, confidence=0.9) for path in frame_paths]


class FakeDescriber:
    def __init__(self, text: str = "A bar chart of quarterly revenue.", error: Exception | None = None) -> None:
        self.text, self.error, self.seen_ocr = text, error, None

    def describe(self, image_path, ocr_text):
        self.seen_ocr = ocr_text
        if self.error:
            raise self.error
        return ImageDescription(text=self.text, model="fake-describer")


class FakeSpeech:
    def transcribe(self, audio_path):
        return TranscriptionResult(
            text="intro body",
            language="en",
            segments=[
                {"start": 0.0, "end": 30.0, "text": "First half minute."},
                {"start": 30.0, "end": 50.0, "text": "Still the first window."},
                {"start": 70.0, "end": 90.0, "text": "A second window."},
            ],
            model="fake-whisper",
        )


def context(ocr=None, speech=None, describer=None) -> ProcessorContext:
    return ProcessorContext(ocr=lambda: ocr or FakeOCR(), speech=lambda: speech, image_describer=lambda: describer)


@pytest.fixture
def small_segments(monkeypatch):
    monkeypatch.setenv("SEGMENT_MAX_CHARACTERS", "60")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def test_markdown_segments_follow_the_heading_path() -> None:
    text = "Intro line.\n\n# Guide\n\nTop.\n\n## Install\n\nRun it.\n\n```\n# not a heading\n```\n\n# Other\n\nEnd."
    segments = split_markdown(text)
    assert [segment.locator.get("section") for segment in segments] == [None, "Guide", "Guide > Install", "Other"]
    assert "# not a heading" in segments[2].text
    assert segments[1].text.startswith("Guide\n\nTop.")


def test_long_sections_split_between_paragraphs(small_segments) -> None:
    paragraphs = "\n\n".join(f"Paragraph number {index} has some words." for index in range(4))
    segments = split_markdown(f"# Notes\n\n{paragraphs}")
    assert len(segments) == 4
    assert [segment.locator["paragraph"] for segment in segments] == [1, 2, 3, 4]
    assert all(segment.locator["section"] == "Notes" for segment in segments)


def test_an_overlong_paragraph_is_cut_at_sentences_then_at_the_limit() -> None:
    parts = pack_paragraphs(["One short. " + "x" * 25 + ". End here."], 20)
    assert all(len(text) <= 20 for text, _ in parts)
    assert "".join(text.replace(" ", "") for text, _ in parts) == "Oneshort." + "x" * 25 + ".Endhere."


def test_pdf_pages_become_page_cited_segments(tmp_path: Path) -> None:
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(make_pdf(["Hybrid retrieval com-\nbines BM25 and dense vectors.", "Second page about rerankers."]))
    ocr = FakeOCR()
    result = process_pdf(pdf, context(ocr=ocr))
    assert [segment.locator for segment in result.segments] == [{"page": 1}, {"page": 2}]
    # Wrapped lines are rejoined and the hyphenation removed.
    assert "combines BM25 and dense vectors" in result.segments[0].text
    assert result.details == {"page_count": 2, "ocr_pages": []}
    assert ocr.calls == 0


def test_a_scanned_page_is_read_with_ocr(tmp_path: Path) -> None:
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(make_pdf(["A page with a real text layer.", ""]))
    result = process_pdf(pdf, context(ocr=FakeOCR("Text recovered by OCR")))
    assert result.segments[1].modality == "ocr"
    assert result.segments[1].locator == {"page": 2}
    assert result.details["ocr_pages"] == [2]
    assert "FakeOCR" in result.models


def test_a_failed_ocr_page_is_a_warning_not_a_failure(tmp_path: Path) -> None:
    pdf = tmp_path / "partial.pdf"
    pdf.write_bytes(make_pdf(["Readable first page.", ""]))
    result = process_pdf(pdf, context(ocr=FakeOCR(error=RuntimeError("tesseract missing"))))
    assert len(result.segments) == 1
    assert result.warnings == ["Page 2 has no text layer and OCR failed: tesseract missing"]


def test_unreadable_pdfs_are_reported(tmp_path: Path) -> None:
    blank = tmp_path / "blank.pdf"
    blank.write_bytes(make_pdf([""]))
    with pytest.raises(ProcessingError, match="even with OCR"):
        process_pdf(blank, context(ocr=FakeOCR(text="")))
    damaged = tmp_path / "damaged.pdf"
    damaged.write_bytes(b"%PDF-1.4\nnot really a pdf")
    with pytest.raises(ProcessingError, match="could not be opened"):
        process_pdf(damaged, context())


def test_word_documents_keep_headings_and_tables(tmp_path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Quarterly review", level=1)
    document.add_paragraph("Revenue grew across regions.")
    document.add_heading("Regions", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Growth"
    table.cell(1, 0).text, table.cell(1, 1).text = "EMEA", "12%"
    path = tmp_path / "review.docx"
    document.save(path)

    result = process_docx(path)
    sections = {segment.locator.get("section"): segment.text for segment in result.segments}
    assert "Revenue grew" in sections["Quarterly review"]
    assert "EMEA | 12%" in sections["Quarterly review > Regions"]
    assert result.models[0] == "python-docx"


def test_the_built_in_html_extractor_keeps_content_and_drops_chrome() -> None:
    html = """
    <html><head><title>T</title><script>var tracking = 1;</script><style>p{}</style></head>
    <body>
      <nav><ul><li>Home<li>About</ul></nav>
      <header>Site banner</header>
      <article><h1>Retrieval guide</h1><p>Hybrid search &amp; reranking.<br>Second line.</p>
      <h2>Steps</h2><ul><li>Index</li><li>Query</li></ul>
      <table><tr><td>k</td><td>60</td></tr></table></article>
      <footer>Copyright</footer>
    </body></html>
    """
    parser = _MarkdownExtractor()
    parser.feed(html)
    parser.close()
    markdown = parser.markdown()
    assert "# Retrieval guide" in markdown and "## Steps" in markdown
    assert "Hybrid search & reranking." in markdown
    assert "- Index" in markdown and "k | 60" in markdown
    for chrome in ("tracking", "Home", "About", "Site banner", "Copyright"):
        assert chrome not in markdown
    # The <title> is the source title already, not body text.
    assert not markdown.startswith("T\n")
    assert html_to_markdown("<p>x</p>")[0].strip()


def test_transcript_windows_cover_spans_of_time() -> None:
    timed = [{"text": f"line {index}", "start": index * 20.0, "duration": 20.0} for index in range(6)]
    windows = window_transcript(timed)
    assert [window.locator for window in windows] == [{"start": 0.0, "end": 60.0}, {"start": 60.0, "end": 120.0}]
    assert windows[0].text == "line 0 line 1 line 2"
    assert all(window.modality == "transcript" for window in windows)


def test_youtube_transcripts_are_windowed(tmp_path: Path) -> None:
    path = tmp_path / "transcript.json"
    path.write_text(json.dumps({"language": "en", "segments": [{"text": "Hello", "start": 1.0, "duration": 2.0}]}))
    result = process_youtube_transcript(path)
    assert result.segments[0].locator == {"start": 1.0, "end": 3.0}
    assert result.duration == 3.0


def test_audio_is_converted_then_transcribed(tmp_path: Path) -> None:
    class FakeMedia:
        def extract_audio(self, path, output_directory):
            return str(tmp_path / "audio.wav"), 90.0

    result = process_audio(tmp_path / "talk.mp3", tmp_path, context(speech=FakeSpeech()), media_processor=FakeMedia())
    assert [segment.locator for segment in result.segments] == [{"start": 0.0, "end": 50.0}, {"start": 70.0, "end": 90.0}]
    assert result.duration == 90.0 and result.models == ["fake-whisper"]
    with pytest.raises(ProcessingError, match="SPEECH_PROVIDER"):
        process_audio(tmp_path / "talk.mp3", tmp_path, context(speech=None), media_processor=FakeMedia())


def test_images_combine_ocr_and_description(tmp_path: Path) -> None:
    image = tmp_path / "chart.png"
    image.write_bytes(b"png")
    describer = FakeDescriber()
    result = process_image(image, context(ocr=FakeOCR("Q1 Q2 Q3"), describer=describer))
    assert [(segment.modality, segment.text) for segment in result.segments] == [
        ("ocr", "Q1 Q2 Q3"),
        ("visual", "A bar chart of quarterly revenue."),
    ]
    # The describer is given the OCR text, so it can read numbers the image makes small.
    assert describer.seen_ocr == "Q1 Q2 Q3"


def test_an_image_survives_one_failed_provider_but_not_two(tmp_path: Path) -> None:
    image = tmp_path / "photo.jpg"
    image.write_bytes(b"jpg")
    result = process_image(image, context(ocr=FakeOCR(error=RuntimeError("no tesseract")), describer=FakeDescriber()))
    assert [segment.modality for segment in result.segments] == ["visual"]
    assert result.warnings == ["OCR failed: no tesseract"]
    with pytest.raises(ProcessingError, match="set VISION_PROVIDER=claude"):
        process_image(image, context(ocr=FakeOCR(text="")))


def test_providers_are_only_built_when_a_processor_needs_them(tmp_path: Path) -> None:
    built: list[str] = []
    lazy = ProcessorContext(
        ocr=lambda: built.append("ocr") or FakeOCR(),
        speech=lambda: built.append("speech"),
        image_describer=lambda: built.append("describer"),
    )
    notes = tmp_path / "notes.md"
    notes.write_text("# Only text\n\nNo models needed.", encoding="utf-8")
    source = Source(kind=SourceKind.document, title="notes", mime_type="text/markdown", media_path=str(notes))
    assert process_source(source, lazy).segments
    assert built == []

    image = tmp_path / "image.png"
    image.write_bytes(b"png")
    process_source(Source(kind=SourceKind.image, title="image", media_path=str(image)), lazy)
    assert built == ["ocr", "describer"]
    lazy.ocr
    assert built == ["ocr", "describer"]


def test_claude_image_descriptions_send_the_image_and_downscale_large_ones(tmp_path: Path) -> None:
    from PIL import Image

    from app.providers.claude import ClaudeImageDescriber

    large = tmp_path / "large.png"
    Image.new("RGB", (3200, 1800), "white").save(large)
    requests: list[dict] = []

    class FakeMessages:
        def create(self, **kwargs):
            requests.append(kwargs)
            return SimpleNamespace(content=[SimpleNamespace(type="text", text="A blank white image.")])

    describer = ClaudeImageDescriber(client=SimpleNamespace(messages=FakeMessages()))
    description = describer.describe(large, "")
    assert description.text == "A blank white image."
    content = requests[0]["messages"][0]["content"]
    assert content[0]["source"]["media_type"] == "image/jpeg"
    assert "(none detected)" in content[1]["text"]
